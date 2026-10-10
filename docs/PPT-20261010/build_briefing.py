#!/usr/bin/env python
"""Build an editable, self-contained meeting deck from existing CPU artifacts.

No model calls, GPU queries, training, sampling, or modification of frozen inputs.
"""
from pathlib import Path
import csv
import hashlib
import json
import math
import sys
import xml.etree.ElementTree as ET

if Path('/tmp/itm-pptdeps').exists():sys.path.insert(0,'/tmp/itm-pptdeps')
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from PIL import Image,ImageDraw,ImageFont
import imageio.v2 as imageio
import imageio_ffmpeg
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR

from briefing_content import slides as slide_content

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
FONT='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
BOLD='/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc'
NAVY='#172D46';TEAL='#187C78';BLUE='#3977AD';GRAY='#7B8794';PURPLE='#8663AC';AMBER='#BC7628'
BG='#F7F9FC';INK='#26384C';MUTED='#697A8C';WHITE='#FFFFFF'
PX=120;W=1600;H=900
fp=FontProperties(fname=FONT)
plt.rcParams.update({'font.family':fp.get_name(),'axes.unicode_minus':False})
COLORS=[GRAY,TEAL,BLUE,PURPLE]
ARMS=['MDM_text_g2.5','original_U100','simple_text_imu','simple_imu_only']
LABELS=['MDM 仅文本','U100 文本＋IMU','简单文本＋IMU','简单 IMU-only']
CHAINS=[[0,2,5,8,11],[0,1,4,7,10],[0,3,6,9,12,15],[9,14,17,19,21],[9,13,16,18,20]]
ASSETS={};LAYOUT=[]


def read(p):return json.loads(Path(p).read_text())
def write(p,x):Path(p).write_text(json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
def bind(path,expected=None):
    p=Path(path);p=p if p.is_absolute() else ROOT/p
    sha=hashlib.sha256(p.read_bytes()).hexdigest()
    if expected and sha!=expected:raise ValueError('Changed input '+str(p))
    ASSETS[str(p.relative_to(ROOT))]=dict(path=str(p.relative_to(ROOT)),sha256=sha,bytes=p.stat().st_size)
    return p


def load_data():
    head=ROOT/'outputs/itm_rework/protocol/head_validation32_evaluation_20261010_v1'
    cand=ROOT/'outputs/itm_rework/protocol/head_primary_evaluation_20261010_v1'
    a=read(bind(head/'aggregates.json'));c=read(bind(cand/'aggregates.json'))
    data=dict(head32={r['arm']:r['metrics'] for r in a},
        candidate8={r['arm']:r['metrics'] for r in c if r['population']=='matched_val8'},
        head32_rows=read(bind(head/'rows.json')),
        candidate_rows=read(bind(cand/'rows.json')),
        comparisons=read(bind(head/'comparison_summary.json')),
        dataset=read(bind(ROOT/'outputs/itm_rework/method/expanded_canonical_inputs_20261009_v1/dataset.json')),
        train_plan=read(bind(ROOT/'docs/itm_rework/protocol/expanded_canonical_training_20261009_v2/plan.json')))
    for folder in [head,cand]:bind(folder/'summary.json')
    if len(data['head32_rows'])!=128:raise ValueError('Full32 comparison required')
    for arm in ARMS:
        rr=[r for r in data['head32_rows'] if r['arm']==arm]
        if len(rr)!=32:raise ValueError('Missing comparison source')
        for k in data['head32'][arm]:
            np.testing.assert_allclose(data['head32'][arm][k],np.mean([r['metrics'][k] for r in rr]),rtol=1e-12,atol=1e-12)
        for r in rr:bind(r['source']['path'],r['source']['sha256'])
    for arm in ['original_U100','head_supervised_U']:
        rr=[r for r in data['candidate_rows'] if r['population']=='matched_val8' and r['arm']==arm]
        if len(rr)!=8:raise ValueError('Candidate requires exactly eight matched validation sources')
        for k in data['candidate8'][arm]:
            np.testing.assert_allclose(data['candidate8'][arm][k],np.mean([r['metrics'][k] for r in rr]),rtol=1e-12,atol=1e-12)
        for r in rr:bind(r['source']['path'],r['source']['sha256'])
    data['head32_vs_MDM_percent']={k:100*(data['head32']['original_U100'][k]/data['head32']['MDM_text_g2.5'][k]-1)
        for k in ['observed_relative_position_mean_m','observed_relative_velocity_RMSE_m_s','root_relative_MPJPE_m',
                  'global_MPJPE_m','observed_relative_jerk_mean_m_s3']}
    return data


def figures(data):
    d=HERE/'figures';d.mkdir(exist_ok=True)
    ms=data['head32']
    panels=[('observed_relative_position_mean_m',100,'头部局部位置误差（cm）'),
        ('observed_relative_velocity_RMSE_m_s',1,'头部局部速度误差（m/s）'),
        ('root_relative_MPJPE_m',100,'全身相对姿态误差（cm）'),
        ('global_MPJPE_m',1,'全身 global 误差（m）')]
    fig,axes=plt.subplots(2,2,figsize=(9.5,7.2),layout='constrained')
    for ax,(k,scale,title) in zip(axes.flat,panels):
        v=[ms[a][k]*scale for a in ARMS];ax.bar(range(4),v,color=COLORS,width=.65)
        ax.set_title(title,fontproperties=fp,fontsize=13)
        ax.set_xticks(range(4),['MDM','U100','简单融合','IMU-only'],fontproperties=fp,fontsize=10)
        ax.set_ylim(0,max(v)*1.2);ax.grid(axis='y',alpha=.2);ax.set_axisbelow(True)
        for i,x in enumerate(v):ax.text(i,x,f'{x:.2f}',ha='center',va='bottom',fontsize=10)
    fig.suptitle('同32个head开发来源 · 单种子／K=1 · 误差越低越好',fontproperties=fp,fontsize=14)
    for ext in ['png','svg','pdf']:fig.savefig(d/('head32_metrics.'+ext),dpi=180)
    p=d/'head32_metrics.svg';p.write_text('\n'.join(line.rstrip() for line in p.read_text().splitlines())+'\n')
    plt.close(fig)
    keys=['observed_relative_position_mean_m','observed_relative_velocity_RMSE_m_s','root_relative_MPJPE_m',
          'global_MPJPE_m','observed_relative_jerk_mean_m_s3']
    labs=['头部位置','头部速度','全身相对姿态','全身global','头部局部jerk']
    vals=[data['head32_vs_MDM_percent'][k] for k in keys]
    fig,ax=plt.subplots(figsize=(7.2,7.2),layout='constrained')
    ax.barh(range(5),vals,color=[TEAL if v<0 else AMBER for v in vals],height=.58)
    ax.set_yticks(range(5),labs,fontproperties=fp,fontsize=13);ax.invert_yaxis();ax.axvline(0,color='#526274',lw=1)
    ax.set_xlim(-42,98);ax.grid(axis='x',alpha=.2);ax.set_axisbelow(True)
    ax.set_xlabel('相对MDM变化（%）；误差／jerk下降为负值',fontproperties=fp,fontsize=11)
    ax.set_title('控制误差下降，但抖动代价仍在',fontproperties=fp,fontsize=15)
    for i,v in enumerate(vals):ax.text(v+(2 if v>=0 else -2),i,f'{v:+.1f}%',ha='left' if v>=0 else 'right',va='center',fontsize=13)
    # Reserve footer space so the caveat does not overlap the x-axis label.
    fig.get_layout_engine().set(rect=(0,.06,1,.94))
    fig.text(.5,.018,'单种子开发比较；jerk低不等于动作自然',ha='center',fontproperties=fp,fontsize=10)
    for ext in ['png','svg','pdf']:fig.savefig(d/('head32_tradeoffs.'+ext),dpi=180)
    p=d/'head32_tradeoffs.svg';p.write_text('\n'.join(line.rstrip() for line in p.read_text().splitlines())+'\n')
    plt.close(fig)
    with (HERE/'tables/head32.csv').open('w',newline='') as f:
        keys=['observed_relative_position_mean_m','observed_relative_velocity_RMSE_m_s','root_relative_MPJPE_m',
              'global_MPJPE_m','root_path_RMSE_m','observed_relative_jerk_mean_m_s3','jerk_absolute_m_s3']
        writer=csv.writer(f,lineterminator='\n');writer.writerow(['method','sources']+keys)
        for a in ARMS:writer.writerow([a,32]+[ms[a][k] for k in keys])
    with (HERE/'tables/head_candidate8.csv').open('w',newline='') as f:
        writer=csv.writer(f,lineterminator='\n');writer.writerow(['method','sources']+keys)
        for a in ['original_U100','head_supervised_U']:writer.writerow([a,8]+[data['candidate8'][a][k] for k in keys])


def font(size,bold=False):return ImageFont.truetype(BOLD if bold else FONT,max(1,int(size*PX/72)))
def wrap(text,ft,width):
    result=[]
    for paragraph in str(text).split('\n'):
        line=''
        for char in paragraph:
            if line and ft.getlength(line+char)>width:result.append(line);line=char
            else:line+=char
        result.append(line)
    return result


def scene(case,frame,size=(1280,600),world=False):
    image=Image.new('RGB',size,WHITE);draw=ImageDraw.Draw(image)
    draw.text((26,15),case['title'],font=font(18,True),fill=NAVY)
    draw.text((26,55),f"{case['split_label']} · head · {frame/20:.2f} / {(case['frames']-1)/20:.2f} s",font=font(11),fill=MUTED)
    n=len(case['panels']);panelw=size[0]/n;origin_y=size[1]-95
    arrays=[np.array(p['motion'],dtype=np.float64) for p in case['panels']]
    def project(x):
        if not world:
            x=x.copy();x[...,0]-=x[...,0,0,None];x[...,2]-=x[...,0,2,None]
        a=.75;b=.35
        depth=-x[...,0]*math.sin(a)+x[...,2]*math.cos(a)
        return np.stack([x[...,0]*math.cos(a)+x[...,2]*math.sin(a),x[...,1]*math.cos(b)-depth*math.sin(b)],-1)
    projs=[project(x) for x in arrays]
    maxx=max(float(np.max(np.abs(x[...,0]))) for x in projs);maxy=max(float(np.max(x[...,1])) for x in projs)
    miny=min(float(np.min(x[...,1])) for x in projs)
    scale=min((panelw-55)/(max(1.6,2*maxx)),(size[1]-180)/max(2.2,maxy-min(0,miny)))
    for i,(p,xy) in enumerate(zip(case['panels'],projs)):
        xc=panelw*(i+.5)
        draw.text((int(i*panelw+18),94),p['label'],font=font(13,True),fill=p['color'])
        draw.line((i*panelw+20,origin_y,(i+1)*panelw-20,origin_y),fill='#D7E1EC',width=2)
        pts=np.stack([xc+xy[frame,:,0]*scale,origin_y-xy[frame,:,1]*scale],-1)
        for chain in CHAINS:
            draw.line([tuple(pts[j]) for j in chain],fill=p['color'],width=5)
        for j,(x,y) in enumerate(pts):draw.ellipse((x-3,y-3,x+3,y+3),fill=p['color'])
        x,y=pts[15];draw.ellipse((x-9,y-9,x+9,y+9),fill='#B44B36',outline=WHITE,width=2)
    footer='显示：各自骨盆XZ居中、保留高度；不展示全局位移。指标未作对齐。' if not world else '显示：原始世界坐标；所有方法共享显示比例。'
    draw.text((25,size[1]-45),footer,font=font(10),fill=MUTED)
    return image


def motions(data):
    cases=[];records={r['motion_id']:r for r in data['dataset']['records']}
    old_sources=read(bind(ROOT/'outputs/itm_rework/protocol/feature_balance_evaluation_20261008_v2/sources.json'))
    old_hashes={r['path']:r['sha256'] for r in old_sources}
    latest=read(bind(ROOT/'outputs/itm_rework/protocol/sensor_relative_evaluation_20261010_v2/rows.json'))
    lookup={(r['arm'],r['motion_id']):r for r in latest if r['stage']=='U100_text2.5' and r['sensor_config']=='head'}
    slots=read(bind(ROOT/'docs/itm_rework/protocol/b0_feature_balance_20261003_v1/sampling_slots.json'))
    actions={'000382':'拍手后放下手臂','000139':'站立／摇摆后行走','000433':'踢腿后返回站立',
             '000027':'出拳后转身','004728':'右腿后举、保持、放下','000506':'交替弓步（失败例）'}
    for mid,title in actions.items():
        r=records[mid];p=bind(r['input']['path'],r['input']['sha256'])
        with np.load(p,allow_pickle=False) as z:gt=z['gt'];acc=z['acceleration'][:,4]
        panels=[dict(label='GT 参考',color=GRAY,motion=gt.tolist())]
        u=lookup['expanded_U',mid];path=bind(u['source']['path'],u['source']['sha256'])
        with np.load(path,allow_pickle=False) as z:pred=z['motion'][0]
        panels.append(dict(label='当前 U100',color=TEAL,motion=pred.tolist()))
        if r['split']=='val':
            m=lookup['MDM_text_g2.5',mid];path=bind(m['source']['path'],m['source']['sha256'])
            with np.load(path,allow_pickle=False) as z:pred=z['motion'][0]
            panels.append(dict(label='MDM 仅文本',color=BLUE,motion=pred.tolist()))
        else:
            s=next(s for s in slots if s['motion_id']==mid and s['sensor_config']=='head' and s['condition']=='paired')
            for label,path,expected,color in [('旧 U（3来源）',s['historical_output']['path'],s['historical_output']['sha256'],BLUE),
                    ('旧 F（损失配重）',str(ROOT/'outputs/itm_rework/method/feature_balance_authorized_20261008_v1/samples'/(s['case_id']+'.npz')),None,PURPLE)]:
                path=bind(path,expected or old_hashes[path])
                with np.load(path,allow_pickle=False) as z:pred=z['motion'][0]
                panels.append(dict(label=label,color=color,motion=pred.tolist()))
        if any(np.asarray(p['motion']).shape!=(r['frames'],22,3) for p in panels):raise ValueError('Whole case required')
        case=dict(motion_id=mid,title=title,caption=r['text'],split=r['split'],
            split_label='开发验证例' if r['split']=='val' else '训练参考例：不作泛化证据',frames=r['frames'],fps=20,panels=panels,
            acceleration=acc.tolist())
        cases.append(case)
        if mid in ['000382','000506','000027','004728']:
            stem=('val_' if r['split']=='val' else 'train_')+mid
            # Representative still is a fixed midpoint; full video retains all frames.
            scene(case,r['frames']//2).save(HERE/'visualizations'/f'{stem}_poster.png')
            writer=imageio_ffmpeg.write_frames(str(HERE/'visualizations'/f'{stem}.mp4'),(1280,600),
                fps=20,codec='libx264',pix_fmt_in='rgb24',pix_fmt_out='yuv420p',macro_block_size=8,output_params=['-crf','21','-movflags','+faststart'])
            writer.send(None)
            for frame in range(r['frames']):writer.send(np.asarray(scene(case,frame)))
            writer.close()
    write(HERE/'visualizations/motion_data.json',dict(cases=cases,selection='All six historical anchors, head only; no output-based sample/seed selection.',
        display_only='Pelvis XZ centering optional; metrics unchanged.',new_model_calls=0,GPU_queries=0))
    player(cases)
    return cases


def player(cases):
    text='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>ITM 离线动作对照</title>
<style>body{font:17px/1.6 system-ui;margin:24px auto;max-width:1380px;padding:0 22px;background:#f7f9fc;color:#172d46}button,select{font-size:16px;padding:8px;margin-right:12px}canvas{width:100%;background:white;border:1px solid #ddd}input{width:90%}.note{color:#697a8c}a{color:#3977ad}</style>
<h1>既有完整动作 · head 配置 · 离线对照</h1><p>六个固定历史锚点，未挑最优输出。训练参考例不代表泛化；旧U/F与当前U100不是单变量消融。<a href="../README.md">材料入口</a></p>
<select id="case"></select><button id="play">暂停</button><button id="mode">切换世界坐标</button><p id="caption"></p><canvas id="canvas" width="1380" height="620"></canvas><input id="time" type="range" min="0" value="0"><span id="clock"></span><p class="note" id="note"></p>
<p><a href="val_000382.mp4">拍手验证例 MP4</a> · <a href="train_000506.mp4">弓步训练失败例 MP4</a> · <a href="train_000027.mp4">出拳转身训练例 MP4</a> · <a href="train_004728.mp4">右腿后举训练例 MP4</a></p>
<script type="application/json" id="data">DATA_PLACEHOLDER</script><script>
const cases=JSON.parse(document.getElementById('data').textContent);let n=0,frame=0,playing=true,world=false,last=0;const select=document.getElementById('case'),cv=document.getElementById('canvas'),ctx=cv.getContext('2d'),slider=document.getElementById('time');
const chains=[[0,2,5,8,11],[0,1,4,7,10],[0,3,6,9,12,15],[9,14,17,19,21],[9,13,16,18,20]];
cases.forEach((c,i)=>{let o=document.createElement('option');o.value=i;o.textContent=c.motion_id+' · '+c.title+' · '+c.split_label;select.appendChild(o)});
function project(x,root){let px=x[0]-(world?0:root[0]),pz=x[2]-(world?0:root[2]),d=-px*Math.sin(.75)+pz*Math.cos(.75);return [px*Math.cos(.75)+pz*Math.sin(.75),x[1]*Math.cos(.35)-d*Math.sin(.35)]}
function draw(){let c=cases[n];slider.max=c.frames-1;slider.value=frame;document.getElementById('caption').textContent=c.caption;document.getElementById('clock').textContent=(frame/20).toFixed(2)+' s';document.getElementById('note').textContent=world?'原始世界坐标；各方法共享显示比例。':'各自骨盆XZ居中、保留高度；指标未作对齐，全局位移不在此模式中展示。';ctx.clearRect(0,0,cv.width,cv.height);let maxx=1,maxy=2.2,miny=0;
c.panels.forEach(p=>p.motion.forEach(f=>f.forEach(x=>{let a=project(x,f[0]);maxx=Math.max(maxx,Math.abs(a[0]));maxy=Math.max(maxy,a[1]);miny=Math.min(miny,a[1])})));let pw=cv.width/c.panels.length,scale=Math.min((pw-55)/(2*maxx),430/(maxy-miny));
c.panels.forEach((p,i)=>{ctx.fillStyle=p.color;ctx.font='bold 21px system-ui';ctx.fillText(p.label,pw*i+20,45);let root=p.motion[frame][0],points=p.motion[frame].map(x=>{let a=project(x,root);return [pw*(i+.5)+a[0]*scale,530-a[1]*scale]});ctx.strokeStyle='#d7e1ec';ctx.lineWidth=2;ctx.beginPath();ctx.moveTo(pw*i+15,530);ctx.lineTo(pw*(i+1)-15,530);ctx.stroke();ctx.strokeStyle=p.color;ctx.lineWidth=5;chains.forEach(ch=>{ctx.beginPath();ch.forEach((j,k)=>{if(k)ctx.lineTo(...points[j]);else ctx.moveTo(...points[j])});ctx.stroke()});ctx.fillStyle='#b44b36';ctx.beginPath();ctx.arc(...points[15],9,0,2*Math.PI);ctx.fill()})}
select.onchange=()=>{n=+select.value;frame=0;draw()};slider.oninput=()=>{frame=+slider.value;playing=false;document.getElementById('play').textContent='播放';draw()};document.getElementById('play').onclick=()=>{playing=!playing;document.getElementById('play').textContent=playing?'暂停':'播放'};document.getElementById('mode').onclick=()=>{world=!world;document.getElementById('mode').textContent=world?'切换骨盆居中':'切换世界坐标';draw()};function loop(t){if(playing&&t-last>=50){frame=(frame+1)%cases[n].frames;last=t;draw()}requestAnimationFrame(loop)}draw();requestAnimationFrame(loop);
</script></html>'''
    (HERE/'visualizations/offline_review.html').write_text(text.replace('DATA_PLACEHOLDER',json.dumps(cases,ensure_ascii=False,separators=(',',':')).replace('</','<\\/')))


class Page:
    def __init__(self,presentation,index,title,subtitle):
        self.slide=presentation.slides.add_slide(presentation.slide_layouts[6])
        self.index=index;self.im=Image.new('RGB',(W,H),BG);self.draw=ImageDraw.Draw(self.im);self.ops=[]
        self.rect(0,0,W/PX,.08,TEAL)
        self.text(title,.6,.29,12.05,.57,30,True,NAVY)
        self.text(subtitle,.62,.94,12,.37,14,False,MUTED)
        self.text('ITM · 导师讨论 · 2026-10-10   |   当前开发实验与研究决策',.62,7.14,11.8,.22,9,False,MUTED)
        self.text(f'{index:02d}',12.4,7.14,.4,.22,9,False,MUTED)

    def rect(self,x,y,w,h,fill=WHITE,stroke=None,radius=0):
        shape=self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            Inches(x),Inches(y),Inches(w),Inches(h))
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(fill[1:])
        if stroke:shape.line.color.rgb=RGBColor.from_string(stroke[1:]);shape.line.width=Pt(1)
        else:shape.line.fill.background()
        coords=(round(x*PX),round(y*PX),round((x+w)*PX),round((y+h)*PX))
        if radius:self.draw.rounded_rectangle(coords,radius=radius,fill=fill,outline=stroke,width=2)
        else:self.draw.rectangle(coords,fill=fill,outline=stroke,width=2)
        self.ops.append(dict(kind='rect',x=x,y=y,w=w,h=h,fill=fill,stroke=stroke))

    def text(self,text,x,y,w,h,size=20,bold=False,color=INK,center=False):
        original=size
        while True:
            ft=font(size,bold);lines=wrap(text,ft,w*PX-10);lead=round(size*PX/72*1.38)
            if lead*len(lines)<=h*PX or size<=9:break
            size-=.5
        if lead*len(lines)>h*PX+2:raise ValueError(f'Text overflow on slide{self.index}: {text}')
        box=self.slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=box.text_frame;tf.clear();tf.word_wrap=False
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        from pptx.enum.text import PP_ALIGN
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if not i else tf.add_paragraph();p.text=line
            p.font.name='Microsoft YaHei';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color[1:])
            p.line_spacing=1.38;p.space_before=p.space_after=Pt(0)
            if center:p.alignment=PP_ALIGN.CENTER
            dx=x*PX+(w*PX-ft.getlength(line))/2 if center else x*PX
            self.draw.text((int(dx),int(y*PX+i*lead)),line,font=ft,fill=color)
        self.ops.append(dict(kind='text',lines=lines,x=x,y=y,w=w,h=h,size=size,bold=bold,color=color,center=center,lead=lead))
        LAYOUT.append(dict(slide=self.index,text=text[:60],requested_size=original,actual_size=size,lines=len(lines),height_fit=True))

    def line(self,x1,y1,x2,y2,color=MUTED,arrow=True,dashed=False):
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        s=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        s.line.color.rgb=RGBColor.from_string(color[1:]);s.line.width=Pt(1.6)
        if dashed:s.line.dash_style=MSO_LINE_DASH_STYLE.DASH
        if arrow:
            tail=OxmlElement('a:tailEnd');tail.set('type','triangle');s.line._get_or_add_ln().append(tail)
        coords=(x1*PX,y1*PX,x2*PX,y2*PX);self.draw.line(coords,fill=color,width=3)
        if arrow:
            a=math.atan2(y2-y1,x2-x1);length=10
            triangle=[(x2*PX,y2*PX),(x2*PX-length*math.cos(a-.5),y2*PX-length*math.sin(a-.5)),
                (x2*PX-length*math.cos(a+.5),y2*PX-length*math.sin(a+.5))]
            self.draw.polygon(triangle,fill=color)
        self.ops.append(dict(kind='line',x1=x1,y1=y1,x2=x2,y2=y2,color=color,arrow=arrow,dashed=dashed))

    def picture(self,path,x,y,w,h,movie=None):
        path=HERE/path;im=Image.open(path).convert('RGB')
        scale=min(w*PX/im.width,h*PX/im.height);nw,nh=im.width*scale,im.height*scale
        xx=x+(w*PX-nw)/2/PX;yy=y+(h*PX-nh)/2/PX
        if movie:self.slide.shapes.add_movie(str(HERE/movie),Inches(xx),Inches(yy),Inches(nw/PX),Inches(nh/PX),
            poster_frame_image=str(path),mime_type='video/mp4')
        else:self.slide.shapes.add_picture(str(path),Inches(xx),Inches(yy),width=Inches(nw/PX),height=Inches(nh/PX))
        self.im.paste(im.resize((round(nw),round(nh)),Image.Resampling.LANCZOS),(round(xx*PX),round(yy*PX)))

    def note(self,text):self.slide.notes_slide.notes_text_frame.text=text

    def save(self):
        p=HERE/'previews'/f'slide_{self.index:02d}.png';self.im.save(p)

    def svg(self,path):
        svg=ET.Element('svg',xmlns='http://www.w3.org/2000/svg',width=str(W),height=str(H),viewBox=f'0 0 {W} {H}')
        ET.SubElement(svg,'rect',x='0',y='0',width=str(W),height=str(H),fill=BG)
        for o in self.ops:
            if o['kind']=='rect':
                ET.SubElement(svg,'rect',x=str(o['x']*PX),y=str(o['y']*PX),width=str(o['w']*PX),height=str(o['h']*PX),
                    fill=o['fill'],stroke=o['stroke'] or 'none')
            elif o['kind']=='line':
                ET.SubElement(svg,'line',x1=str(o['x1']*PX),y1=str(o['y1']*PX),x2=str(o['x2']*PX),y2=str(o['y2']*PX),
                    stroke=o['color'],**{'stroke-width':'3','stroke-dasharray':'8 6' if o['dashed'] else 'none'})
                if o['arrow']:
                    a=math.atan2(o['y2']-o['y1'],o['x2']-o['x1']);xx=o['x2']*PX;yy=o['y2']*PX
                    points=[(xx,yy),(xx-10*math.cos(a-.5),yy-10*math.sin(a-.5)),(xx-10*math.cos(a+.5),yy-10*math.sin(a+.5))]
                    ET.SubElement(svg,'polygon',points=' '.join(f'{x},{y}' for x,y in points),fill=o['color'])
            else:
                for i,line in enumerate(o['lines']):
                    xx=(o['x']+o['w']/2)*PX if o['center'] else o['x']*PX
                    t=ET.SubElement(svg,'text',x=str(xx),y=str(o['y']*PX+i*o['lead']+o['size']*PX/72),fill=o['color'],
                        **{'font-family':'Microsoft YaHei,Noto Sans CJK SC,sans-serif','font-size':str(o['size']*PX/72),
                           'font-weight':'bold' if o['bold'] else 'normal','text-anchor':'middle' if o['center'] else 'start'})
                    t.text=line
        ET.ElementTree(svg).write(path,encoding='utf-8',xml_declaration=True)


def cards(page,items):
    n=len(items);gap=.24;width=(12.1-gap*(n-1))/n
    for i,(title,body) in enumerate(items):
        x=.62+i*(width+gap)
        page.rect(x,1.55,width,4.52,WHITE,'#E0E7F0',10)
        page.rect(x,1.55,width,.08,[TEAL,BLUE,PURPLE][i%3])
        page.text(title,x+.20,1.84,width-.4,.77,23,True,[TEAL,BLUE,PURPLE][i%3])
        page.text(body,x+.20,2.78,width-.4,2.95,20)


def node(page,label,x,y,w,h,color=BLUE):
    page.rect(x,y,w,h,WHITE,color,8);page.text(label,x+.07,y+.06,w-.14,h-.10,17,True,color,True)


def architecture(page):
    node(page,'文本描述',.65,1.70,2.15,.68)
    node(page,'冻结 CLIP',3.20,1.70,1.80,.68,NAVY)
    page.line(2.80,2.04,3.20,2.04);page.line(5.00,2.04,5.70,2.04)
    page.text('噪声 x_t ＋ 扩散时间步 t',5.85,1.38,3.3,.35,14,False,MUTED)
    page.rect(5.70,1.82,3.88,2.07,'#EDF2FA',NAVY,9)
    page.text('冻结 MDM Transformer ×8',5.92,1.98,3.40,.40,17,True,NAVY)
    for i in range(3):
        node(page,'Transformer → ＋ A_l(C)',5.96,2.48+i*.41,3.3,.35,NAVY)
    page.line(9.58,2.8,10.03,2.8)
    node(page,'预测 x0\nDDPM 迭代 1000 步\n20 Hz 全身动作',10.05,1.86,2.57,1.92,TEAL)
    node(page,'头部 IMU 观测\na(t) ＋ R(t)',.65,3.32,2.15,.80,AMBER)
    node(page,'物理静态参照\na=0，R=首帧朝向',.65,4.55,2.15,.87,AMBER)
    node(page,'E(I)\n时序编码器',3.20,3.32,1.80,.80,TEAL)
    node(page,'E(I_ref)\n时序编码器',3.20,4.55,1.80,.87,TEAL)
    page.text('共享权重 · 秒级时间编码',2.98,4.16,2.35,.33,12,False,TEAL,True)
    page.line(2.80,3.73,3.20,3.73,AMBER);page.line(2.80,4.98,3.20,4.98,AMBER)
    page.line(5.00,3.73,5.55,4.06,TEAL);page.line(5.00,4.98,5.55,4.65,TEAL)
    node(page,'−',5.57,3.98,.68,.78,TEAL)
    node(page,'C = 编码差',6.62,4.05,2.08,.69,TEAL)
    page.line(6.25,4.38,6.62,4.38,TEAL);page.line(7.69,4.05,7.69,3.90,TEAL)
    page.text('逐层线性残差适配器\n权重零初始化，bias固定0',6.56,4.90,3.22,.57,13,False,TEAL)
    node(page,'GT 动作（仅训练）',10.05,4.60,2.57,.62,GRAY)
    page.text('标准化 x0 MSE',9.72,5.38,2.9,.45,14,True,MUTED,True)
    page.line(11.30,5.22,11.30,5.40,GRAY,dashed=True)
    page.line(12.39,3.78,12.39,5.59,TEAL,dashed=True)
    page.line(12.39,5.59,12.13,5.59,TEAL,dashed=True)
    page.text('GT不进入IMU／推理输入',9.32,5.90,3.45,.27,11,False,MUTED)
    page.svg(HERE/'figures/Figure1_method.svg')
    page.im.save(HERE/'figures/Figure1_method.png')


def pipeline(page):
    node(page,'AMASS 源动作\n＋HumanML 文本',.70,2.55,2.25,1.12,NAVY)
    node(page,'SMPL FK\n同源坐标映射',3.35,2.55,2.25,1.12,NAVY)
    page.line(2.95,3.12,3.35,3.12)
    node(page,'目标动作\nHumanML263／22关节',6.04,1.88,2.58,1.10,BLUE)
    node(page,'虚拟头部 IMU\n顶点加速度／segment朝向',6.04,3.57,2.58,1.18,AMBER)
    page.line(5.60,3.10,6.04,2.46,BLUE);page.line(5.60,3.10,6.04,4.11,AMBER)
    node(page,'统一20Hz真实时间\n加速度线性／朝向SLERP',9.08,3.57,3.06,1.18,AMBER)
    page.line(8.62,4.16,9.08,4.16,AMBER)
    node(page,'配对训练与评价\n按来源隔离 train／val',9.08,1.88,3.06,1.10,TEAL)
    page.line(8.62,2.43,9.08,2.43,BLUE);page.line(10.61,3.57,10.61,2.99,TEAL)
    page.text('加速度：世界线性加速度（不加重力）\n朝向：segment → world；缓存30Hz → 模型20Hz\n离线canonical映射与真实设备在线定标是两件事',1.05,5.28,11.5,.91,17,False,MUTED)
    page.svg(HERE/'figures/data_pipeline.svg');page.im.save(HERE/'figures/data_pipeline.png')


def tables(data):
    m=data['head32'];c=data['candidate8']
    result={}
    result['baselines']=(['模型','当前状态','输入','主要用途'],[
        ['MDM 文本基座','32val已完成','文本','量化IMU条件增量'],
        ['简单时序融合','100epoch／32val','文本＋head IMU','检验复杂控制是否必要'],
        ['简单IMU-only','100epoch／32val','head IMU','重建代理；不是官方模型'],
        ['官方IMUPoser h','待适配、训练','head IMU','主要重建强基线'],
        ['Ego4o','待资产与接口','按公开分支核查','后续多模态生成／重建对照'],
        ['HY-Motion／MotionLab','历史探索','协议不同','不混入当前共同主表']], [2.4,2.6,2.2,4.9])
    result['head32']=(['方法','头部位置\ncm ↓','头部速度\nm/s ↓','全身相对\ncm ↓','global\nm ↓','头部jerk\nm/s³'],[
        [LABELS[i],f"{m[a]['observed_relative_position_mean_m']*100:.2f}",f"{m[a]['observed_relative_velocity_RMSE_m_s']:.3f}",
         f"{m[a]['root_relative_MPJPE_m']*100:.2f}",f"{m[a]['global_MPJPE_m']:.3f}",f"{m[a]['observed_relative_jerk_mean_m_s3']:.2f}"]
        for i,a in enumerate(ARMS)],[3.1,1.7,1.8,1.8,1.5,2.2])
    keys=['observed_relative_position_mean_m','observed_relative_velocity_RMSE_m_s','global_MPJPE_m','root_relative_MPJPE_m','observed_relative_jerk_mean_m_s3']
    result['candidate']=(['同8val','头部位置\ncm','头部速度\nm/s','global\nm','全身相对\ncm','头部jerk\nm/s³'],[
        ['原U100']+[f'{c["original_U100"][k]*s:.3f}' for k,s in zip(keys,[100,1,1,100,1])],
        ['局部监督候选']+[f'{c["head_supervised_U"][k]*s:.3f}' for k,s in zip(keys,[100,1,1,100,1])],
        ['相对原U变化']+[f'{100*(c["head_supervised_U"][k]/c["original_U100"][k]-1):+.1f}%' for k in keys]],
        [2.7,1.9,1.9,1.7,1.9,2.0])
    result['completion']=(['交付','状态','还缺什么'],[
        ['基本任务／完整生成','工程闭环已完成','行为与真实设备有效性仍待验证'],
        ['最低开发对照','4方法×32head来源','官方IMUPoser／共同语义指标'],
        ['头部辅助损失候选','完整比较；未采用','不继续扫系数；保留负结果'],
        ['核心结构消融','尚未完成','参照差分／时序设计同预算重训'],
        ['文本与IMU互补','尚未充分量化','去条件／文本编辑／共同语义评测'],
        ['正式测试／多种子','尚未完成','最终模型固定后，来源隔离test'],
        ['论文Method／Experiments','正文尚未更新','先固定任务与实验，再冻结主张']], [3.0,3.4,5.7])
    result['roadmap']=(['交付顺序','具体工作','完成判据'],[
        ['1 官方head基线','修正12／60维接口；canonical旋转\n20Hz与SMPL FK；一次完整训练','有来源／权重／预测可核验的32val表'],
        ['2 联合条件证据','下载并绑定语义评测资产\n共同GT／文本／长度检查，条件贡献','说明文本是否补充未观测部位语义'],
        ['3 核心消融与冻结','两项保留设计同预算重训\n扩大数据，正式test与多种子','同一最终模型＋公平主表＋负例保留']], [2.1,5.3,4.7])
    result['full_metrics']=(['方法','全身速度\nm/s ↓','全身jerk\nm/s³','根轨迹\nm ↓'],[
        [LABELS[i],f"{m[a]['velocity_joint_proxy_RMSE_m_s']:.3f}",f"{m[a]['jerk_absolute_m_s3']:.2f}",
         f"{m[a]['root_path_RMSE_m']:.3f}"] for i,a in enumerate(ARMS)],[4.0,2.7,2.7,2.7])
    return result


def draw_table(page,table):
    headers,rows,widths=table
    widths=np.asarray(widths)*12.1/sum(widths)
    x=.62;y=1.60;hh=.84;rh=min(.65,4.20/max(1,len(rows)))
    for i,(heading,w) in enumerate(zip(headers,widths)):
        page.rect(x,y,float(w),hh,NAVY)
        page.text(heading,x+.12,y+.08,float(w)-.24,hh-.15,17,True,WHITE)
        x+=w
    for j,row in enumerate(rows):
        x=.62;yy=y+hh+j*rh
        highlight=row[0] in ['U100 文本＋IMU','原U100']
        for value,w in zip(row,widths):
            page.rect(x,yy,float(w),rh,'#E5F2EF' if highlight else WHITE,'#E0E7F0')
            page.text(value,x+.12,yy+.07,float(w)-.24,rh-.12,17,highlight,TEAL if highlight else INK)
            x+=w


def build_deck(data):
    presentation=Presentation();presentation.slide_width=Inches(W/PX);presentation.slide_height=Inches(H/PX)
    presentation.core_properties.title='ITM：文本＋头部IMU全身动作生成 · 2026-10-10导师汇报'
    presentation.core_properties.subject='实际开发进展、方法与实验、研究路线复盘'
    presentation.core_properties.author='ITM research project'
    content=slide_content(data);all_tables=tables(data);pages=[]
    md=['# ITM 导师汇报逐页讲稿','', '建议主稿约15–20分钟；01–22页为主稿，23–25页为附录。','',
        '表格和图来自同32来源开发比较；单种子/K=1，未认证语义、临床或最终论文优越性。','']
    for index,s in enumerate(content,1):
        page=Page(presentation,index,s['title'],s['subtitle']);kind=s['kind']
        if kind in ['cover','cards']:cards(page,s['cards'])
        elif kind=='architecture':architecture(page)
        elif kind=='pipeline':pipeline(page)
        elif kind=='table':draw_table(page,all_tables[s['table']])
        elif kind=='figures':
            page.picture(s['images'][0],.55,1.50,7.15,4.75);page.picture(s['images'][1],7.72,1.50,5.1,4.75)
        elif kind=='motion':page.picture(s['image'],.62,1.45,12.1,4.7,movie=s['video'])
        page.rect(.62,6.45,12.1,.57,'#EAF1F8',None,8)
        page.text(s['takeaway'],.81,6.52,11.71,.40,18,True,NAVY)
        page.note(s['notes']);page.save();pages.append(page)
        md += [f'## {index:02d} {s["title"]}','',s['subtitle'],'']
        if 'cards' in s:
            for title,body in s['cards']:md += [f'**{title}**','']+['- '+line for line in body.split('\n')]+['']
        elif 'table' in s:
            headers,rows,_=all_tables[s['table']]
            md+=['| '+' | '.join(h.replace('\n',' ') for h in headers)+' |','| '+' | '.join('---' for _ in headers)+' |']
            md+=['| '+' | '.join(str(x).replace('\n','；') for x in row)+' |' for row in rows]+['']
        elif 'images' in s:md += [f'![{s["title"]}]({p})' for p in s['images']]+['']
        elif 'image' in s:md += [f'![{s["title"]}]({s["image"]})',f'[完整动作MP4]({s["video"]})','']
        md += ['**本页结论：**'+s['takeaway'],'','**建议讲述：**'+s['notes'],'']
    ppt=HERE/'ITM_导师汇报_20261010.pptx';presentation.save(ppt)
    (HERE/'汇报逐页讲稿.md').write_text('\n'.join(md))
    pages[0].im.save(HERE/'ITM_导师汇报_20261010.pdf',save_all=True,append_images=[p.im for p in pages[1:]],resolution=120)
    thumbs=Image.new('RGB',(1600,math.ceil(len(pages)/4)*245),WHITE)
    for i,p in enumerate(pages):
        im=p.im.resize((390,220),Image.Resampling.LANCZOS);thumbs.paste(im,((i%4)*400,(i//4)*245))
    thumbs.save(HERE/'previews/contact_sheet.png')
    html='<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>ITM汇报预览</title><style>body{background:#e5eaf0;font:16px system-ui;max-width:1280px;margin:auto}img{width:100%;background:white;margin:12px 0}a{color:#187c78}</style><h1>25页汇报预览</h1><p>程序绘制的同布局预览；不是桌面PowerPoint截图。<a href="../visualizations/offline_review.html">离线动作对照</a></p>'
    html+=''.join(f'<h2>{i:02d} {s["title"]}</h2><img src="slide_{i:02d}.png">' for i,s in enumerate(content,1))
    (HERE/'previews/index.html').write_text(html)
    write(HERE/'sources/layout_checks.json',LAYOUT)
    return content


def main():
    for folder in ['figures','tables','visualizations','previews','sources']:
        (HERE/folder).mkdir(exist_ok=True)
    data=load_data()
    print('Verified 128 existing head predictions and their 32-source aggregates.',flush=True)
    figures(data)
    cases=motions(data)
    print('Rendered six complete historical cases; four full MP4s.',flush=True)
    content=build_deck(data)
    # Bind implementation and reports, without altering experiment artifacts.
    for source in [
        'src/itm/models/mdm_imu_control.py',
        'docs/itm_rework/CURRENT_STATE.md',
        'docs/itm_rework/protocol/HEAD_VALIDATION32_EVALUATION_20261010.md',
        'docs/itm_rework/protocol/HEAD_PRIMARY_CANDIDATE_EVALUATION_20261010.md',
        'docs/itm_rework/HEAD_IMUPOSER_INTERFACE_ADDENDUM_20261010.md',
        'docs/itm_rework/HEAD_SEMANTIC_EVALUATOR_READINESS_20261010.md',
    ]:bind(ROOT/source)
    write(HERE/'sources/source_manifest.json',dict(
        repository_head='4da3b5b',report_date='2026-10-10',
        inputs=list(ASSETS.values()),new_model_calls=0,GPU_queries=0,
        selection='Four methods, same 32 head validation sources; six fixed historical motion anchors.',
        display='Skeleton pelvis XZ centering only in visualization; all reported metrics unchanged.',
        slides=len(content),main_slides=22,appendix_slides=3,
        motion_cases=[dict(motion_id=c['motion_id'],frames=c['frames'],split=c['split']) for c in cases]))
    write(HERE/'tables/metrics_snapshot.json',dict(head32=data['head32'],candidate8=data['candidate8'],
        head32_vs_MDM_percent=data['head32_vs_MDM_percent']))
    print('Saved editable PPTX, preview PDF, speaker notes, charts and input manifest.',flush=True)


if __name__=='__main__':main()
