#!/usr/bin/env python
"""Verify the meeting package's data bindings, structure, media and offline player."""
from pathlib import Path
import hashlib,json,re,subprocess,sys,zipfile,posixpath
import xml.etree.ElementTree as ET
if Path('/tmp/itm-pptdeps').exists():sys.path.insert(0,'/tmp/itm-pptdeps')
import numpy as np
from PIL import Image
from pptx import Presentation
import imageio_ffmpeg

HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[1]
def read(p):return json.loads(p.read_text())
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
    manifest=read(HERE/'sources/source_manifest.json')
    for r in manifest['inputs']:
        p=ROOT/r['path'];assert p.stat().st_size==r['bytes'] and sha(p)==r['sha256'],r['path']
    snapshot=read(HERE/'tables/metrics_snapshot.json')
    rows=read(ROOT/'outputs/itm_rework/protocol/head_validation32_evaluation_20261010_v1/rows.json')
    same=set()
    scalar_checks=0
    for arm,metrics in snapshot['head32'].items():
        rr=[r for r in rows if r['arm']==arm];ids={r['motion_id'] for r in rr}
        assert len(rr)==len(ids)==32
        assert not same or same==ids;same=ids
        for key,value in metrics.items():
            np.testing.assert_allclose(value,np.mean([r['metrics'][key] for r in rr]),rtol=1e-12,atol=1e-12)
            scalar_checks+=1
    pp=Presentation(HERE/'ITM_导师汇报_20261010.pptx');assert len(pp.slides)==25
    shape_count=0;notes_count=0
    for slide in pp.slides:
        assert slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()
        notes_count+=1
        for s in slide.shapes:
            assert s.left>=-2 and s.top>=-2
            assert s.left+s.width<=pp.slide_width+2 and s.top+s.height<=pp.slide_height+2
            shape_count+=1
    with zipfile.ZipFile(HERE/'ITM_导师汇报_20261010.pptx') as z:
        assert z.testzip() is None
        xmls=[p for p in z.namelist() if p.endswith('.xml') or p.endswith('.rels')]
        for p in xmls:ET.fromstring(z.read(p))
        embedded=[p for p in z.namelist() if p.endswith('.mp4')];assert len(embedded)==2
        expected={sha(HERE/'visualizations'/p) for p in ['val_000382.mp4','train_000506.mp4']}
        assert {hashlib.sha256(z.read(p)).hexdigest() for p in embedded}==expected
        for p in z.namelist():
            if not p.endswith('.rels'):continue
            base=posixpath.dirname(posixpath.dirname(p))
            for r in ET.fromstring(z.read(p)):
                if r.attrib.get('TargetMode')=='External':continue
                dest=posixpath.normpath(posixpath.join(base,r.attrib['Target'])).lstrip('/')
                assert dest in z.namelist(),(p,dest)
    motion=read(HERE/'visualizations/motion_data.json')['cases'];assert len(motion)==6
    videos=[]
    for c in motion:
        for p in c['panels']:
            a=np.asarray(p['motion']);assert a.shape==(c['frames'],22,3) and np.isfinite(a).all()
        stem=('val_' if c['split']=='val' else 'train_')+c['motion_id']
        p=HERE/'visualizations'/f'{stem}.mp4'
        if not p.exists():continue
        reader=imageio_ffmpeg.read_frames(str(p));meta=next(reader)
        n=sum(1 for _ in reader);reader.close()
        assert n==c['frames'] and meta['size']==(1280,600) and meta['fps']==20
        videos.append(dict(path=str(p.relative_to(HERE)),frames=n,fps=meta['fps'],whole_clip=True))
    assert len(videos)==4
    for i in range(1,26):
        with Image.open(HERE/'previews'/f'slide_{i:02d}.png') as im:assert im.size==(1600,900)
    pdf=(HERE/'ITM_导师汇报_20261010.pdf').read_bytes()
    assert len(re.findall(rb'/Type\s*/Page\b',pdf))==25
    for f in (HERE/'figures').glob('*.svg'):ET.parse(f)
    checked_links=0
    for p in HERE.rglob('*.md'):
        for link in re.findall(r'\]\(([^)]+)\)',p.read_text()):
            if '://' in link or link.startswith('#'):continue
            target=(p.parent/link.split('#')[0]).resolve()
            assert target.exists() or target==HERE/'sources/verification.json',(p.name,link)
            checked_links+=1
    js=json.loads(subprocess.run(['node',str(HERE/'verify_player.js'),str(HERE/'visualizations/offline_review.html')],
        check=True,capture_output=True,text=True,timeout=60).stdout)
    report=dict(status='passed',input_SHA256_checked=len(manifest['inputs']),aggregate_scalar_checks=scalar_checks,
        slides=25,speaker_notes=notes_count,shapes_checked=shape_count,embedded_full_videos=2,
        standalone_full_videos=videos,PDF_pages=25,local_markdown_links_checked=checked_links,
        offline_player=js,desktop_Office_rendering='not available; program-rendered layout previews inspected',
        experiment_mutations=0,new_model_calls=0,GPU_queries=0)
    (HERE/'sources/verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
